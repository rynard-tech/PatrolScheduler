from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Callable, Iterable

from openpyxl import load_workbook

from .models import Employee


@dataclass(frozen=True)
class ImportIssue:
    row: int
    code: str
    message: str
    severity: str = "ERROR"


@dataclass
class ImportPreview:
    records: list[dict[str, Any]] = field(default_factory=list)
    issues: list[ImportIssue] = field(default_factory=list)
    matched: int = 0
    new: int = 0

    @property
    def requires_review(self) -> bool:
        return any(issue.severity == "ERROR" for issue in self.issues)


class SpreadsheetImporter:
    """Schema-mapped workbook reader; source columns are never assumed."""

    def preview(
        self,
        path: str | Path,
        column_mapping: dict[str, str],
        required_fields: Iterable[str],
        sheet_name: str | None = None,
        transforms: dict[str, Callable[[Any], Any]] | None = None,
    ) -> ImportPreview:
        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook[sheet_name] if sheet_name else workbook.active
        rows = sheet.iter_rows(values_only=True)
        try:
            headers = [
                str(value).strip() if value is not None else "" for value in next(rows)
            ]
        except StopIteration:
            return ImportPreview(
                issues=[ImportIssue(1, "EMPTY_SHEET", "Workbook sheet is empty")]
            )
        positions = {header: index for index, header in enumerate(headers)}
        preview = ImportPreview()
        for target, source in column_mapping.items():
            if source not in positions:
                preview.issues.append(
                    ImportIssue(
                        1,
                        "MISSING_COLUMN",
                        f"Mapped column '{source}' for '{target}' is absent",
                    )
                )
        if preview.requires_review:
            return preview
        transforms = transforms or {}
        for row_number, values in enumerate(rows, start=2):
            if not any(value is not None for value in values):
                continue
            record = {
                target: values[positions[source]]
                for target, source in column_mapping.items()
            }
            for field_name, transform in transforms.items():
                if field_name in record and record[field_name] is not None:
                    try:
                        record[field_name] = transform(record[field_name])
                    except (TypeError, ValueError) as exc:
                        preview.issues.append(
                            ImportIssue(
                                row_number, "INVALID_VALUE", f"{field_name}: {exc}"
                            )
                        )
            for required in required_fields:
                if record.get(required) in (None, ""):
                    preview.issues.append(
                        ImportIssue(
                            row_number,
                            "MISSING_VALUE",
                            f"Required field '{required}' is blank",
                        )
                    )
            preview.records.append(record)
        return preview


class IdentityMatcher:
    """Only IDs and unique exact normalized names auto-match; fuzzy matches require review."""

    def __init__(self, employees: Iterable[Employee]):
        self.by_number: dict[str, list[Employee]] = {}
        self.by_name: dict[str, list[Employee]] = {}
        for employee in employees:
            if employee.employee_number:
                self.by_number.setdefault(employee.employee_number.strip(), []).append(
                    employee
                )
            self.by_name.setdefault(
                self.normalize_name(employee.display_name), []
            ).append(employee)

    @staticmethod
    def normalize_name(name: str) -> str:
        return " ".join(name.casefold().split())

    def match(
        self, row: int, employee_number: str | None, name: str | None
    ) -> tuple[Employee | None, ImportIssue | None]:
        candidates = (
            self.by_number.get(employee_number.strip(), []) if employee_number else []
        )
        if not candidates and name:
            candidates = self.by_name.get(self.normalize_name(name), [])
        if len(candidates) == 1:
            return candidates[0], None
        if len(candidates) > 1:
            return None, ImportIssue(
                row,
                "AMBIGUOUS_IDENTITY",
                "Multiple exact employee matches require admin review",
            )
        return None, ImportIssue(
            row,
            "UNRESOLVED_IDENTITY",
            "No exact employee match; admin confirmation is required",
        )
